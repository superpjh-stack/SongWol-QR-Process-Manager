#!/usr/bin/env python3
"""IMS 엑셀 다운로드 파일 프로파일링 (spec §12.4 1단계).

사용법:
    python migration/profile.py <xlsx 경로> [--max-key-cols 3]

시트마다 다음을 출력한다.
- 헤더(1행) 컬럼 목록과 데이터 건수
- 컬럼별 결측률 (빈 셀 / None 비율)
- 중복 키 후보: 단일 컬럼 및 최대 N개 컬럼 조합 중 값이 유일한(중복 0) 것
- 값이 유일하지 않지만 중복률이 낮은 컬럼은 "거의 유일" 로 표시

의존: openpyxl (backend/.venv 에 설치됨).
"""

from __future__ import annotations

import argparse
import sys
from collections import Counter
from itertools import combinations
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _is_missing(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def profile_sheet(rows: list[tuple[Any, ...]], max_key_cols: int) -> dict[str, Any]:
    if not rows:
        return {"columns": [], "count": 0, "missing": {}, "unique_keys": [], "near_unique": []}

    header = [str(h) if h is not None else f"col{i + 1}" for i, h in enumerate(rows[0])]
    data = [r for r in rows[1:] if not all(_is_missing(c) for c in r)]
    n = len(data)

    missing: dict[str, float] = {}
    for i, col in enumerate(header):
        miss = sum(1 for r in data if i >= len(r) or _is_missing(r[i]))
        missing[col] = (miss / n) if n else 0.0

    unique_keys: list[tuple[str, ...]] = []
    near_unique: list[tuple[str, float]] = []
    candidate_idx = [i for i, col in enumerate(header) if missing[col] < 0.05]

    for k in range(1, min(max_key_cols, len(candidate_idx)) + 1):
        for combo in combinations(candidate_idx, k):
            # 상위 조합이 이미 유일하면 그 확장은 볼 필요 없다
            if any(set(u) <= {header[i] for i in combo} for u in unique_keys):
                continue
            keys = [tuple(r[i] if i < len(r) else None for i in combo) for r in data]
            c = Counter(keys)
            dup = sum(v - 1 for v in c.values() if v > 1)
            if n and dup == 0:
                unique_keys.append(tuple(header[i] for i in combo))
            elif k == 1 and n and dup / n < 0.02:
                near_unique.append((header[combo[0]], dup / n))

    return {
        "columns": header,
        "count": n,
        "missing": missing,
        "unique_keys": unique_keys,
        "near_unique": near_unique,
    }


def profile_workbook(path: Path, max_key_cols: int) -> None:
    wb = load_workbook(path, read_only=True, data_only=True)
    print(f"# {path}")
    for ws in wb.worksheets:
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
        p = profile_sheet(rows, max_key_cols)
        print(f"\n## 시트: {ws.title}  (건수 {p['count']}, 컬럼 {len(p['columns'])})")
        print("컬럼 | 결측률")
        for col in p["columns"]:
            print(f"  {col} | {p['missing'][col]:.1%}")
        print("중복 없는 키 후보:")
        if p["unique_keys"]:
            for k in p["unique_keys"]:
                print(f"  {' + '.join(k)}")
        else:
            print("  (없음)")
        if p["near_unique"]:
            print("거의 유일 (중복률 <2%):")
            for col, rate in p["near_unique"]:
                print(f"  {col} | 중복률 {rate:.2%}")
    wb.close()


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx", type=Path, help="IMS 다운로드 엑셀 파일")
    ap.add_argument("--max-key-cols", type=int, default=3, help="중복 키 후보 조합 최대 컬럼 수")
    args = ap.parse_args(argv)
    if not args.xlsx.exists():
        print(f"파일 없음: {args.xlsx}", file=sys.stderr)
        return 2
    profile_workbook(args.xlsx, args.max_key_cols)
    return 0


if __name__ == "__main__":
    sys.exit(main())
