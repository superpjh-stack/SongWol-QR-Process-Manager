"""S0-6 엑셀 일괄 등록.

템플릿 → 100행 xlsx(오류 3행) → preview → commit → 건수·migration_batch · 두 번 commit 409
· 중복 병합(SKIP/UPDATE) · 배치 조회/discard · 권한 · 크기/확장자.
"""

import io

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from tests.conftest import ensure_item_group, headers_for, uniq

API = "/api/v1"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def make_xlsx(headers: list[str], rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def customer_rows(tag: str, n: int = 100) -> tuple[list[str], list[list[object]]]:
    headers = [
        "code",
        "name",
        "contact_name",
        "phone",
        "email",
        "legacy_id",
        "addr_label",
        "addr_address1",
    ]
    rows: list[list[object]] = []
    for i in range(n):
        rows.append(
            [
                f"IC{tag}-{i:03d}",
                f"거래처 {tag} {i}",
                "담당",
                f"031-{i:04d}",
                None,
                f"L{tag}{i}",
                "본사",
                f"주소 {i}",
            ]
        )
    # 오류 3행: 필수 누락(name) · 길이 초과(code 21자) · 파일 내 중복 코드
    rows[10][1] = None
    rows[20][0] = "X" * 21
    rows[30][0] = rows[0][0]
    return headers, rows


async def upload(
    client: AsyncClient,
    headers: dict[str, str],
    entity: str,
    data: bytes,
    name: str = "u.xlsx",
    source: str | None = None,
):
    form = {"entity": entity}
    if source:
        form["source"] = source
    return await client.post(
        f"{API}/master/import/preview",
        headers=headers,
        data=form,
        files={"file": (name, data, XLSX_MIME)},
    )


@pytest.mark.asyncio
async def test_template_download(client: AsyncClient, admin_headers: dict[str, str]) -> None:
    for entity, first in (("customer", "code"), ("item", "code"), ("stock", "item_code")):
        res = await client.get(
            f"{API}/master/import/template", headers=admin_headers, params={"entity": entity}
        )
        assert res.status_code == 200, res.text
        assert res.headers["content-type"].startswith(XLSX_MIME)
        assert f"import_template_{entity}.xlsx" in res.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(res.content))
        assert wb.sheetnames == ["data", "설명"]
        ws = wb["data"]
        header = [c.value for c in ws[1]]
        assert header[0] == first and ws.max_row == 2  # 헤더 + 예시 1행
    res = await client.get(
        f"{API}/master/import/template", headers=admin_headers, params={"entity": "wo"}
    )
    assert res.status_code == 422


@pytest.mark.asyncio
async def test_customer_import_100_rows_preview_commit_twice(
    client: AsyncClient, admin_headers: dict[str, str]
) -> None:
    tag = uniq("")
    headers, rows = customer_rows(tag)
    data = make_xlsx(headers, rows)
    res = await upload(client, admin_headers, "customer", data, name=f"customers_{tag}.xlsx")
    assert res.status_code == 201, res.text
    pv = res.json()
    assert pv["entity"] == "customer" and pv["source"] == "IMS_XLS"
    assert pv["row_count"] == 100 and pv["valid"] == 97
    assert {e["row"] for e in pv["errors"]} == {12, 22, 32}  # 엑셀 행 번호 (헤더 1행)
    assert {e["col"] for e in pv["errors"]} == {"name", "code"}
    assert pv["duplicates"] == [] and len(pv["rows_sample"]) == 10
    batch_id = pv["batch_id"]
    # 오류 포함 commit → 409 IMPORT_HAS_ERRORS
    res = await client.post(
        f"{API}/master/import/{batch_id}/commit",
        headers=admin_headers,
        json={"merge_policy": "SKIP"},
    )
    assert res.status_code == 409 and res.json()["code"] == "IMPORT_HAS_ERRORS"
    # skip_invalid → 오류 행 제외 적재
    res = await client.post(
        f"{API}/master/import/{batch_id}/commit",
        headers=admin_headers,
        json={"merge_policy": "SKIP", "skip_invalid": True},
    )
    assert res.status_code == 200, res.text
    assert res.json() == {
        "batch_id": batch_id,
        "loaded": 97,
        "merged": 0,
        "skipped": 0,
        "failed": 3,
    }
    # migration_batch 확인 (row_count_src = loaded + merged + failed)
    from app.db.models.master import Customer, CustomerAddress
    from app.db.models.ops import MigrationBatch, MigrationMap
    from app.db.session import SessionLocal

    async with SessionLocal() as s:
        b = await s.get(MigrationBatch, batch_id)
        assert b is not None and b.status == "LOADED" and b.merge_policy == "SKIP"
        assert b.row_count_src == 100 and b.row_count_loaded == 97 and b.row_count_merged == 0
        assert (
            len(b.errors) == 3
            and b.source_file == f"customers_{tag}.xlsx"
            and len(b.source_hash) == 64
        )
        n = (
            (await s.execute(select(Customer).where(Customer.code.like(f"IC{tag}-%"))))
            .scalars()
            .all()
        )
        assert len(n) == 97
        maps = (
            (await s.execute(select(MigrationMap).where(MigrationMap.batch_id == batch_id)))
            .scalars()
            .all()
        )
        assert len(maps) == 97 and all(m.entity == "customer" for m in maps)
        addr = (
            (await s.execute(select(CustomerAddress).where(CustomerAddress.customer_id == n[0].id)))
            .scalars()
            .all()
        )
        assert len(addr) == 1 and addr[0].is_default is True and addr[0].label == "본사"
    # 두 번째 commit → 409 STATE_CONFLICT
    res = await client.post(
        f"{API}/master/import/{batch_id}/commit",
        headers=admin_headers,
        json={"merge_policy": "SKIP", "skip_invalid": True},
    )
    assert res.status_code == 409 and res.json()["code"] == "STATE_CONFLICT"
    # 같은 파일 재업로드 → duplicates 97 (CODE) → UPDATE 병합
    res = await upload(client, admin_headers, "customer", data)
    pv2 = res.json()
    assert len(pv2["duplicates"]) == 97 and {d["reason"] for d in pv2["duplicates"]} == {"CODE"}
    res = await client.post(
        f"{API}/master/import/{pv2['batch_id']}/commit",
        headers=admin_headers,
        json={"merge_policy": "UPDATE", "skip_invalid": True},
    )
    assert res.status_code == 200 and res.json()["merged"] == 97 and res.json()["loaded"] == 0
    # 세 번째: SKIP → skipped 97
    res = await upload(client, admin_headers, "customer", data)
    res = await client.post(
        f"{API}/master/import/{res.json()['batch_id']}/commit",
        headers=admin_headers,
        json={"merge_policy": "SKIP", "skip_invalid": True},
    )
    assert res.json()["skipped"] == 97 and res.json()["loaded"] == 0
    # 배치 목록 · 상세(errors · duplicates · maps Page)
    res = await client.get(
        f"{API}/migration/batches",
        headers=admin_headers,
        params={"entity": "customer", "status": "LOADED"},
    )
    assert res.status_code == 200 and res.json()["total"] >= 3
    assert res.json()["items"][0]["created_by"]["login_id"]
    res = await client.get(
        f"{API}/migration/batches/{batch_id}", headers=admin_headers, params={"size": 5}
    )
    assert res.status_code == 200
    d = res.json()
    assert len(d["errors"]) == 3 and d["maps"]["total"] == 97 and len(d["maps"]["items"]) == 5


@pytest.mark.asyncio
async def test_item_import_with_item_group_check_and_name_phone_dup(
    client: AsyncClient, admin_headers: dict[str, str]
) -> None:
    tag = uniq("")
    group = await ensure_item_group(f"IG{tag}")
    headers = ["code", "name", "item_group", "spec", "weight_g", "qty_tolerance_pct"]
    rows: list[list[object]] = [
        [f"II{tag}-{i:03d}", f"품목 {i}", group, "40×80", 150, 3.0] for i in range(100)
    ]
    rows[5][2] = "NO_GROUP"  # 활성 품목군 아님
    rows[6][4] = -1  # 음수 중량
    rows[7][5] = 99  # 허용오차 초과
    res = await upload(client, admin_headers, "item", make_xlsx(headers, rows))
    assert res.status_code == 201, res.text
    pv = res.json()
    assert pv["valid"] == 97 and {e["col"] for e in pv["errors"]} == {
        "item_group",
        "weight_g",
        "qty_tolerance_pct",
    }
    res = await client.post(
        f"{API}/master/import/{pv['batch_id']}/commit",
        headers=admin_headers,
        json={"merge_policy": "SKIP", "skip_invalid": True},
    )
    assert res.status_code == 200 and res.json()["loaded"] == 97
    # 거래처 상호+전화 일치 → NAME_PHONE 중복 후보 (코드는 다름)
    c_headers = ["code", "name", "phone"]
    first = make_xlsx(c_headers, [[f"NP{tag}-A", f"상사{tag}", "010-0000"]])
    pv1 = (await upload(client, admin_headers, "customer", first)).json()
    await client.post(
        f"{API}/master/import/{pv1['batch_id']}/commit",
        headers=admin_headers,
        json={"merge_policy": "SKIP"},
    )
    second = make_xlsx(c_headers, [[f"NP{tag}-B", f"상사{tag}", "010-0000"]])
    pv2 = (await upload(client, admin_headers, "customer", second)).json()
    assert pv2["duplicates"] == [{"row": 2, "existing_code": f"NP{tag}-A", "reason": "NAME_PHONE"}]
    # 미리보기만 한 배치 discard → ROLLED_BACK, 이후 commit 불가
    res = await client.post(
        f"{API}/migration/batches/{pv2['batch_id']}/discard", headers=admin_headers
    )
    assert res.status_code == 200 and res.json()["status"] == "ROLLED_BACK"
    res = await client.post(
        f"{API}/master/import/{pv2['batch_id']}/commit",
        headers=admin_headers,
        json={"merge_policy": "SKIP"},
    )
    assert res.status_code == 409


@pytest.mark.asyncio
async def test_import_rejections_and_permissions(
    client: AsyncClient, admin_headers: dict[str, str]
) -> None:
    # 필수 컬럼 없는 템플릿 → 422 BAD_TEMPLATE · 엑셀 아님 → 422 BAD_FILE_TYPE · 확장자 · 크기 413
    res = await upload(client, admin_headers, "customer", make_xlsx(["foo"], [["x"]]))
    assert res.status_code == 422 and res.json()["code"] == "BAD_TEMPLATE"
    res = await upload(client, admin_headers, "customer", b"not an excel")
    assert res.status_code == 422 and res.json()["code"] == "BAD_FILE_TYPE"
    res = await upload(
        client, admin_headers, "customer", make_xlsx(["code", "name"], [["a", "b"]]), name="u.csv"
    )
    assert res.status_code == 422 and res.json()["code"] == "BAD_FILE_TYPE"
    res = await upload(client, admin_headers, "customer", b"0" * (5 * 1024 * 1024 + 1))
    assert res.status_code == 413 and res.json()["code"] == "FILE_TOO_LARGE"
    res = await upload(
        client, admin_headers, "stock", make_xlsx(["item_code", "qty_on_hand"], [["a", 1]])
    )
    assert res.status_code == 422 and res.json()["code"] == "ENTITY_NOT_SUPPORTED"
    # 권한: MANAGER 는 customer/item 임포트 불가, SALES 는 가능 · stock 은 ADMIN/MANAGER
    manager = await headers_for(client, "t_manager", "MANAGER")
    sales = await headers_for(client, "t_sales", "SALES")
    assert (
        await client.get(
            f"{API}/master/import/template", headers=manager, params={"entity": "customer"}
        )
    ).status_code == 403
    assert (
        await client.get(
            f"{API}/master/import/template", headers=sales, params={"entity": "customer"}
        )
    ).status_code == 200
    assert (
        await client.get(f"{API}/master/import/template", headers=sales, params={"entity": "stock"})
    ).status_code == 403
    assert (
        await client.get(
            f"{API}/master/import/template", headers=manager, params={"entity": "stock"}
        )
    ).status_code == 200
    assert (await client.get(f"{API}/migration/batches", headers=manager)).status_code == 200
    assert (await client.get(f"{API}/migration/batches", headers=sales)).status_code == 403
    res = await upload(
        client, sales, "customer", make_xlsx(["code", "name"], [[uniq("SC"), "영업등록"]])
    )
    assert res.status_code == 201
    assert (
        await client.post(
            f"{API}/migration/batches/{res.json()['batch_id']}/discard", headers=manager
        )
    ).status_code == 403
