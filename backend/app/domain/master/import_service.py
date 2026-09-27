"""엑셀 일괄 등록 (A1-11 · spec §12 · api-contract §7.2 · §13.3 admin #12/#22/#23/#24).

흐름: 템플릿 다운로드 → preview(검증 · 중복 후보 · migration_batch PREVIEW) → commit(적재) → 배치
조회/discard.
- 템플릿: 시트 1 ``data`` = 헤더 + 예시 1행, 시트 2 ``설명`` = 컬럼 설명. 첫 행 헤더, ``.xlsx``.
- 상한: 5MB (413 FILE_TOO_LARGE) · 5,000행 (422 TOO_MANY_ROWS). 필수 컬럼 누락 422 BAD_TEMPLATE.
- 원본 파일은 ``import_dir/{batch_id}.xlsx`` 로 보관(spec §12.2 해시와 함께). commit 은 파일을
  다시 읽어 재검증한다
  (미리보기 뒤 DB 가 바뀌었을 수 있다).
- 중복 후보: code 일치 = CODE, 거래처 상호+전화 일치 = NAME_PHONE. commit ``merge_policy``
  SKIP/UPDATE. NAME_PHONE 은 SKIP 이 기본 처리, UPDATE 는 명시 선택 시 기존 코드 행 갱신
  (``migration_map.note='MERGED_INTO:{existing_code}'``, D31).
- 첫 열이 ``#`` 로 시작하는 행(템플릿 예시 ``#예시``)은 무시하고 ``ignored`` 로 센다 (D29).
- preview 와 commit 은 같은 ``validate_rows`` 를 쓴다(활성 택배사·품목군 포함, D30). commit
  재검증에서 오류가 나면 배치는 PREVIEW 유지 + errors 갱신 → 재commit 가능. 적재 중 업무 오류는 422
  VALIDATION_ERROR(PREVIEW 유지), FAILED 는 적재 중 예외(DB 오류 등)에만.
- ``errors>0`` 이고 ``skip_invalid`` 없으면 409 IMPORT_HAS_ERRORS. 같은 entity 의 동시 commit 은
  advisory lock → 423 LOCKED.
- entity=stock 은 S5(마이그레이션 실사 반영) — 템플릿만 제공, preview 는 422 ENTITY_NOT_SUPPORTED.
"""

from __future__ import annotations

import hashlib
import io
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.v1.schemas import master as S
from app.api.v1.schemas.common import Page
from app.core.config import get_settings
from app.core.errors import ApiError, not_found, state_conflict, validation
from app.db.models.master import AppUser, Carrier, Customer, CustomerAddress, Item, ItemGroup
from app.db.models.ops import MigrationBatch, MigrationMap
from app.domain.common.listing import PageParams, paginate, parse_sort
from app.domain.master.admin_service import user_summary

DATA_SHEET = "data"
DOC_SHEET = "설명"
SAMPLE_ROWS = 10


@dataclass(frozen=True, slots=True)
class Col:
    name: str
    required: bool = False
    max_len: int | None = None
    kind: str = "str"  # str | int | decimal
    min_val: float | None = None
    max_val: float | None = None
    desc: str = ""
    example: str = ""


# admin #22 — 템플릿 컬럼 (``*`` 필수)
COLUMNS: dict[str, list[Col]] = {
    "customer": [
        Col("code", True, 20, desc="거래처 코드 (IMS 코드 유지)", example="C-0001"),
        Col("name", True, 100, desc="상호", example="한빛상사"),
        Col("contact_name", False, 50, desc="담당자", example="김담당"),
        Col("phone", False, 30, desc="연락처", example="02-123-4567"),
        Col("email", False, 100, desc="이메일", example="order@example.com"),
        Col(
            "default_carrier",
            False,
            20,
            desc="기본 택배사 (택배사 마스터가 있으면 그 코드)",
            example="",
        ),
        Col("legacy_id", False, 50, desc="IMS 행 식별값", example="IMS-1001"),
        Col(
            "addr_label",
            False,
            50,
            desc="배송지명 (주소 열이 있으면 기본 배송지 1건 생성)",
            example="본사",
        ),
        Col("addr_receiver", False, 50, desc="수령인", example="김담당"),
        Col("addr_phone", False, 30, desc="배송지 연락처", example="010-1234-5678"),
        Col("addr_postal_code", False, 10, desc="우편번호", example="06236"),
        Col("addr_address1", False, 200, desc="주소", example="서울 강남구 테헤란로 1"),
        Col("addr_address2", False, 200, desc="상세주소", example="3층"),
    ],
    "item": [
        Col("code", True, 30, desc="품목 코드", example="TW-40-WH"),
        Col("name", True, 100, desc="품목명", example="40수 세면타월 화이트"),
        Col("item_group", True, 30, desc="품목군 코드 (활성 품목군이어야 함)", example="TOWEL_40"),
        Col("spec", False, 50, desc="규격", example="40×80"),
        Col("color", False, 30, desc="색상", example="화이트"),
        Col("weight_g", False, kind="int", min_val=0, desc="중량(g)", example="150"),
        Col("vendor_item_code", False, 50, desc="협력업체 품번", example="V-001"),
        Col("vendor_barcode", False, 64, desc="협력업체 바코드", example="8801234567890"),
        Col(
            "qty_tolerance_pct",
            False,
            kind="decimal",
            min_val=0,
            max_val=50,
            desc="수량 허용오차(%) 기본 3.0",
            example="3.0",
        ),
        Col("legacy_id", False, 50, desc="IMS 행 식별값", example="IMS-2001"),
    ],
    "stock": [
        Col("item_code", True, 30, desc="품목 코드", example="TW-40-WH"),
        Col("qty_on_hand", True, kind="int", min_val=0, desc="실사 재고 수량", example="1200"),
        Col("note", False, 200, desc="비고", example=""),
    ],
}
DEFAULT_SOURCE: dict[str, str] = {"customer": "IMS_XLS", "item": "IMS_XLS", "stock": "COUNT"}
SUPPORTED_COMMIT = frozenset({"customer", "item"})

BATCH_SORT: dict[str, Any] = {"created_at": MigrationBatch.created_at}


@dataclass(slots=True)
class ParsedRow:
    row: int
    values: dict[str, Any]
    errors: list[dict[str, Any]] = field(default_factory=list)
    duplicate: dict[str, Any] | None = None

    @property
    def valid(self) -> bool:
        return not self.errors


@dataclass(slots=True)
class Validation:
    rows: list[ParsedRow]
    errors: list[dict[str, Any]]
    duplicates: list[dict[str, Any]]

    @property
    def valid_rows(self) -> list[ParsedRow]:
        return [r for r in self.rows if r.valid]


# ======================================================================
# 템플릿
# ======================================================================
def build_template(entity: str) -> bytes:
    cols = COLUMNS.get(entity)
    if cols is None:
        raise validation(["query", "entity"], f"알 수 없는 entity: {entity}")
    wb = Workbook()
    ws = wb.active
    ws.title = DATA_SHEET
    ws.append([c.name for c in cols])
    example = [c.example for c in cols]
    example[0] = f"#예시 {example[0]}"  # D29: '#' 시작 행은 preview 가 무시한다
    ws.append(example)
    doc = wb.create_sheet(DOC_SHEET)
    doc.append(["컬럼", "필수", "형식", "최대 길이", "설명"])
    for c in cols:
        rng = ""
        if c.min_val is not None or c.max_val is not None:
            lo = c.min_val if c.min_val is not None else ""
            hi = c.max_val if c.max_val is not None else ""
            rng = f" ({lo}~{hi})"
        doc.append([c.name, "Y" if c.required else "", c.kind + rng, c.max_len or "", c.desc])
    doc.append([])
    doc.append(
        [
            "규칙",
            "첫 행 헤더 · 시트 1개(data) · .xlsx · 최대 5MB · 5,000행 · "
            "첫 열이 '#' 로 시작하는 행(2행 예시)은 무시된다",
        ]
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ======================================================================
# 파싱 · 검증
# ======================================================================
def _cell_text(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def parse_workbook(data: bytes, entity: str) -> tuple[list[ParsedRow], int]:
    """→ (rows, ignored). 첫 열이 '#' 로 시작하는 행은 ignored 로 센다 (D29)."""
    cols = COLUMNS[entity]
    ignored = 0
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # openpyxl 은 zip/xml 오류를 여러 예외로 낸다
        raise validation(["body", "file"], "엑셀(.xlsx) 형식이 아닙니다", "BAD_FILE_TYPE") from e
    ws = wb[DATA_SHEET] if DATA_SHEET in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header_row = next(it, None)
    if header_row is None:
        raise validation(["body", "file"], "헤더 행이 없습니다", "BAD_TEMPLATE")
    header = [(_cell_text(h) or "").strip() for h in header_row]
    idx = {name: i for i, name in enumerate(header) if name}
    missing = [c.name for c in cols if c.required and c.name not in idx]
    if missing:
        err = validation(
            ["body", "file"],
            f"필수 컬럼이 없습니다: {', '.join(missing)} (템플릿을 다시 받으세요)",
            "BAD_TEMPLATE",
        )
        err.detail.append({"missing_columns": missing})
        raise err
    rows: list[ParsedRow] = []
    max_rows = get_settings().import_max_rows
    for excel_row, raw in enumerate(it, start=2):
        values = {
            c.name: (
                _cell_text(raw[idx[c.name]])
                if idx.get(c.name) is not None and idx[c.name] < len(raw)
                else None
            )
            for c in cols
        }
        if all(v is None for v in values.values()):
            continue
        first = values.get(cols[0].name)
        if isinstance(first, str) and first.startswith("#"):
            ignored += 1
            continue
        rows.append(ParsedRow(row=excel_row, values=values))
        if len(rows) > max_rows:
            raise validation(["body", "file"], f"행이 {max_rows}개를 넘습니다", "TOO_MANY_ROWS")
    return rows, ignored


def _validate_cell(r: ParsedRow, c: Col) -> None:
    v = r.values.get(c.name)
    if v is None:
        if c.required:
            r.errors.append({"row": r.row, "col": c.name, "msg": "필수 값입니다"})
        return
    if c.kind == "str":
        if c.max_len is not None and len(v) > c.max_len:
            r.errors.append({"row": r.row, "col": c.name, "msg": f"{c.max_len}자를 넘습니다"})
        return
    try:
        num = Decimal(v) if c.kind == "decimal" else Decimal(int(v))
    except (InvalidOperation, ValueError):
        r.errors.append({"row": r.row, "col": c.name, "msg": "숫자여야 합니다"})
        return
    if c.min_val is not None and num < Decimal(str(c.min_val)):
        r.errors.append({"row": r.row, "col": c.name, "msg": f"{c.min_val} 이상이어야 합니다"})
    if c.max_val is not None and num > Decimal(str(c.max_val)):
        r.errors.append({"row": r.row, "col": c.name, "msg": f"{c.max_val} 이하여야 합니다"})
    r.values[c.name] = int(num) if c.kind == "int" else num


async def validate_rows(
    session: AsyncSession, entity: str, rows: Sequence[ParsedRow]
) -> Validation:
    cols = COLUMNS[entity]
    seen: dict[str, int] = {}
    for r in rows:
        for key in ("code", "item_group", "item_code"):
            val = r.values.get(key)
            if isinstance(val, str):
                r.values[key] = val.strip().upper()  # D28 코드 정규화 (vendor_barcode 는 원문 유지)
        for c in cols:
            _validate_cell(r, c)
        code = r.values.get("code")
        if isinstance(code, str):
            if code in seen:
                r.errors.append(
                    {"row": r.row, "col": "code", "msg": f"파일 내 중복 코드 (행 {seen[code]})"}
                )
            else:
                seen[code] = r.row
    if entity == "item":
        groups = set(
            (await session.execute(select(ItemGroup.code).where(ItemGroup.active.is_(True))))
            .scalars()
            .all()
        )
        for r in rows:
            g = r.values.get("item_group")
            if g is not None and g not in groups:
                r.errors.append(
                    {"row": r.row, "col": "item_group", "msg": f"활성 품목군이 아닙니다: {g}"}
                )
    if entity == "customer":
        # D30: commit(check_carrier) 과 같은 규칙 — 활성 택배사가 있으면 그 코드만
        carriers = set(
            (await session.execute(select(Carrier.code).where(Carrier.active.is_(True))))
            .scalars()
            .all()
        )
        for r in rows:
            dc = r.values.get("default_carrier")
            if carriers and dc is not None and dc not in carriers:
                r.errors.append(
                    {
                        "row": r.row,
                        "col": "default_carrier",
                        "msg": f"등록된 택배사가 아닙니다: {dc}",
                    }
                )
            addr_cols = [k for k, v in r.values.items() if k.startswith("addr_") and v is not None]
            if addr_cols and r.values.get("addr_address1") is None:
                r.errors.append(
                    {
                        "row": r.row,
                        "col": "addr_address1",
                        "msg": "주소 열을 쓰려면 addr_address1 이 필요합니다",
                    }
                )

    # 중복 후보 (DB)
    codes = [r.values["code"] for r in rows if isinstance(r.values.get("code"), str)]
    model: type[Customer] | type[Item] = Customer if entity == "customer" else Item
    existing_by_code: dict[str, str] = {}
    if codes:
        found = (
            (await session.execute(select(model.code).where(model.code.in_(codes)))).scalars().all()
        )
        existing_by_code = {c: c for c in found}
    name_phone: dict[tuple[str, str], str] = {}
    if entity == "customer":
        pairs = [
            (r.values["name"], r.values["phone"])
            for r in rows
            if r.values.get("name") and r.values.get("phone")
        ]
        if pairs:
            q = select(Customer.code, Customer.name, Customer.phone).where(
                Customer.name.in_({p[0] for p in pairs}), Customer.phone.in_({p[1] for p in pairs})
            )
            for code, name, phone in (await session.execute(q)).all():
                if phone is not None:
                    name_phone[(name, phone)] = code
    for r in rows:
        if r.errors:
            continue  # 오류 행은 적재되지 않으므로 중복 후보에서 제외
        code = r.values.get("code")
        if isinstance(code, str) and code in existing_by_code:
            r.duplicate = {"row": r.row, "existing_code": code, "reason": "CODE"}
        elif entity == "customer":
            name_v, phone_v = r.values.get("name"), r.values.get("phone")
            if isinstance(name_v, str) and isinstance(phone_v, str):
                hit = name_phone.get((name_v, phone_v))
                if hit is not None:
                    r.duplicate = {"row": r.row, "existing_code": hit, "reason": "NAME_PHONE"}
    errors = [e for r in rows for e in r.errors]
    duplicates = [r.duplicate for r in rows if r.duplicate is not None]
    return Validation(rows=list(rows), errors=errors, duplicates=duplicates)


# ======================================================================
# 배치
# ======================================================================
def _batch_path(batch_id: int) -> Path:
    return Path(get_settings().import_dir) / f"{batch_id}.xlsx"


async def batch_out(session: AsyncSession, b: MigrationBatch) -> S.MigrationBatch:
    user = await session.get(AppUser, b.created_by)
    assert user is not None
    return S.MigrationBatch(
        id=b.id,
        source=b.source,
        entity=b.entity,
        source_file=b.source_file,
        source_hash=b.source_hash,
        extracted_at=b.extracted_at,
        row_count_src=b.row_count_src,
        row_count_loaded=b.row_count_loaded,
        row_count_merged=b.row_count_merged,
        row_count_skipped=b.row_count_skipped,
        row_count_failed=b.row_count_failed,
        row_count_ignored=b.row_count_ignored,
        status=b.status,
        merge_policy=b.merge_policy,
        created_by=user_summary(user),
        created_at=b.created_at,
    )


async def get_batch(session: AsyncSession, batch_id: int) -> MigrationBatch:
    b = await session.get(MigrationBatch, batch_id)
    if b is None:
        raise not_found("BATCH_NOT_FOUND", "마이그레이션 배치", batch_id)
    return b


async def preview(
    session: AsyncSession,
    *,
    entity: str,
    source: str | None,
    filename: str,
    data: bytes,
    user_id: int,
) -> S.ImportPreview:
    if entity not in COLUMNS:
        raise validation(["body", "entity"], f"알 수 없는 entity: {entity}")
    if entity not in SUPPORTED_COMMIT:
        raise validation(
            ["body", "entity"],
            "기초재고(stock) 임포트는 S5 마이그레이션 단계에서 제공합니다",
            "ENTITY_NOT_SUPPORTED",
        )
    if not filename.lower().endswith(".xlsx"):
        raise validation(["body", "file"], "확장자는 .xlsx 여야 합니다", "BAD_FILE_TYPE")
    settings = get_settings()
    if len(data) > settings.import_max_bytes:
        raise ApiError(
            413,
            "FILE_TOO_LARGE",
            f"파일이 {settings.import_max_bytes // (1024 * 1024)}MB 를 넘습니다",
        )
    src = source or DEFAULT_SOURCE[entity]
    if src not in ("IMS_XLS", "COUNT"):
        raise validation(["body", "source"], "source 는 IMS_XLS 또는 COUNT")

    rows, ignored = parse_workbook(data, entity)
    result = await validate_rows(session, entity, rows)
    batch = MigrationBatch(
        source=src,
        entity=entity,
        source_file=filename,
        source_hash=hashlib.sha256(data).hexdigest(),
        extracted_at=datetime.now(UTC),
        row_count_src=len(rows),
        row_count_failed=len({e["row"] for e in result.errors}),
        row_count_ignored=ignored,
        errors=result.errors,
        duplicates=result.duplicates,
        created_by=user_id,
    )
    session.add(batch)
    await session.flush()
    path = _batch_path(batch.id)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    await session.commit()
    sample = [dict(r.values) for r in result.valid_rows[:SAMPLE_ROWS]]
    return S.ImportPreview(
        batch_id=batch.id,
        entity=entity,
        source=src,
        row_count=len(rows),
        valid=len(result.valid_rows),
        ignored=ignored,
        errors=[S.ImportError(**e) for e in result.errors],
        duplicates=[S.ImportDuplicate(**d) for d in result.duplicates],
        rows_sample=_jsonable(sample),
    )


def _jsonable(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [{k: (float(v) if isinstance(v, Decimal) else v) for k, v in r.items()} for r in rows]


def _customer_fields(v: dict[str, Any]) -> dict[str, Any]:
    return {
        k: v.get(k)
        for k in ("name", "contact_name", "phone", "email", "default_carrier", "legacy_id")
    }


def _item_fields(v: dict[str, Any]) -> dict[str, Any]:
    out: dict[str, Any] = {
        k: v.get(k)
        for k in (
            "name",
            "item_group",
            "spec",
            "color",
            "weight_g",
            "vendor_item_code",
            "vendor_barcode",
            "legacy_id",
        )
    }
    if v.get("qty_tolerance_pct") is not None:
        out["qty_tolerance_pct"] = v["qty_tolerance_pct"]
    return out


async def _load_rows(
    session: AsyncSession, batch: MigrationBatch, result: Validation, merge_policy: str
) -> tuple[int, int, int]:
    """(loaded, merged, skipped). 호출자 트랜잭션 안."""
    entity = batch.entity
    loaded = merged = skipped = 0
    for r in result.valid_rows:
        v = r.values
        if r.duplicate is not None:
            if merge_policy == "SKIP":
                skipped += 1
                continue
            existing_code = r.duplicate["existing_code"]
            if entity == "customer":
                cust = (
                    await session.execute(select(Customer).where(Customer.code == existing_code))
                ).scalar_one()
                for k, val in _customer_fields(v).items():
                    if val is not None:
                        setattr(cust, k, val)
                new_id = cust.id
            else:
                item = (
                    await session.execute(select(Item).where(Item.code == existing_code))
                ).scalar_one()
                for k, val in _item_fields(v).items():
                    if val is not None:
                        setattr(item, k, val)
                new_id = item.id
            merged += 1
            note = (
                f"MERGED_INTO:{existing_code}"
                if r.duplicate["reason"] == "NAME_PHONE"
                else f"merge CODE row {r.row}"
            )
        else:
            if entity == "customer":
                cust = Customer(code=v["code"], **_customer_fields(v))
                session.add(cust)
                await session.flush()
                if v.get("addr_address1"):
                    session.add(
                        CustomerAddress(
                            customer_id=cust.id,
                            label=v.get("addr_label") or "기본",
                            receiver=v.get("addr_receiver"),
                            phone=v.get("addr_phone"),
                            postal_code=v.get("addr_postal_code"),
                            address1=v["addr_address1"],
                            address2=v.get("addr_address2"),
                            is_default=True,
                        )
                    )
                new_id = cust.id
            else:
                item = Item(code=v["code"], **_item_fields(v))
                session.add(item)
                await session.flush()
                new_id = item.id
            loaded += 1
            note = None
        legacy = v.get("legacy_id")
        if legacy:
            already = (
                await session.execute(
                    select(MigrationMap.id).where(
                        MigrationMap.entity == entity, MigrationMap.legacy_id == legacy
                    )
                )
            ).first()
            if already is None:
                session.add(
                    MigrationMap(
                        batch_id=batch.id, entity=entity, legacy_id=legacy, new_id=new_id, note=note
                    )
                )
    return loaded, merged, skipped


async def commit(
    session: AsyncSession, batch_id: int, body: S.ImportCommitRequest
) -> S.ImportResult:
    batch = await get_batch(session, batch_id)
    if batch.status != "PREVIEW":
        raise state_conflict(
            f"배치 {batch_id} 은(는) {batch.status} 상태라 적재할 수 없습니다 (PREVIEW 만 가능)"
        )
    if batch.entity not in SUPPORTED_COMMIT:
        raise validation(["path", "batch_id"], "기초재고(stock) 적재는 S5", "ENTITY_NOT_SUPPORTED")
    lock_row = await session.execute(
        text("SELECT pg_try_advisory_xact_lock(hashtext(:k))"), {"k": f"import:{batch.entity}"}
    )
    locked = bool(lock_row.scalar_one())
    if not locked:
        raise ApiError(
            423, "LOCKED", f"{batch.entity} 적재가 진행 중입니다 — 잠시 후 다시 시도하세요"
        )
    path = _batch_path(batch.id)
    if not path.is_file():
        raise state_conflict(f"배치 {batch_id} 의 원본 파일이 없습니다 — 다시 업로드하세요")
    data = path.read_bytes()
    if hashlib.sha256(data).hexdigest() != batch.source_hash:
        raise state_conflict(f"배치 {batch_id} 의 원본 파일 해시가 다릅니다 — 다시 업로드하세요")

    rows, ignored = parse_workbook(data, batch.entity)
    result = await validate_rows(session, batch.entity, rows)
    # D30: 재검증 결과는 항상 배치에 남긴다 (미리보기 이후 DB 가 바뀌었을 수 있다).
    # 상태는 PREVIEW 유지.
    batch.errors = result.errors
    batch.duplicates = result.duplicates
    batch.row_count_failed = len({e["row"] for e in result.errors})
    batch.row_count_ignored = ignored
    if result.errors and not body.skip_invalid:
        await session.commit()
        raise ApiError(
            409,
            "IMPORT_HAS_ERRORS",
            f"오류 행이 {len({e['row'] for e in result.errors})}건 있습니다. "
            "파일을 고치거나 skip_invalid=true 로 제외 적재하세요",
            result.errors,
        )
    failed = len({e["row"] for e in result.errors})
    try:
        loaded, merged, skipped = await _load_rows(session, batch, result, body.merge_policy)
        batch.row_count_loaded = loaded
        batch.row_count_merged = merged
        batch.row_count_skipped = skipped
        batch.row_count_failed = failed
        batch.merge_policy = body.merge_policy
        batch.status = "LOADED"
        await session.commit()
    except ApiError as e:
        # 적재 중 업무 검증 오류 → PREVIEW 유지 + errors 갱신, 422 VALIDATION_ERROR (D30).
        # 재commit 허용
        await session.rollback()
        kept = await get_batch(session, batch_id)
        kept.errors = result.errors + [{"row": 0, "col": "", "msg": e.message}]
        kept.duplicates = result.duplicates
        await session.commit()
        raise ApiError(
            422, "VALIDATION_ERROR", f"적재 중 검증 오류: {e.message}", kept.errors
        ) from e
    except Exception:
        # 적재 중 예외(DB 오류 등)에만 FAILED — 재commit 불가, 새 preview
        await session.rollback()
        failed_batch = await get_batch(session, batch_id)
        failed_batch.status = "FAILED"
        await session.commit()
        raise
    return S.ImportResult(
        batch_id=batch.id, loaded=loaded, merged=merged, skipped=skipped, failed=failed
    )


async def discard(session: AsyncSession, batch_id: int) -> S.MigrationBatch:
    batch = await get_batch(session, batch_id)
    if batch.status != "PREVIEW":
        raise state_conflict(f"배치 {batch_id} 은(는) {batch.status} 상태라 취소할 수 없습니다")
    batch.status = "ROLLED_BACK"
    await session.commit()
    path = _batch_path(batch.id)
    if path.is_file():
        path.unlink()
    return await batch_out(session, batch)


async def list_batches(
    session: AsyncSession,
    params: PageParams,
    *,
    entity: str | None,
    status: str | None,
    source: str | None,
    sort: str | None,
) -> tuple[list[S.MigrationBatch], int]:
    stmt = select(MigrationBatch)
    if entity:
        stmt = stmt.where(MigrationBatch.entity == entity)
    if status:
        stmt = stmt.where(MigrationBatch.status == status)
    if source:
        stmt = stmt.where(MigrationBatch.source == source)
    stmt = stmt.order_by(*parse_sort(sort, BATCH_SORT, default="-created_at"))
    rows, total = await paginate(session, stmt, params)
    return [await batch_out(session, b) for b in rows], total


async def batch_detail(
    session: AsyncSession, batch_id: int, maps_params: PageParams
) -> S.MigrationBatchDetail:
    b = await get_batch(session, batch_id)
    base = await batch_out(session, b)
    stmt = select(MigrationMap).where(MigrationMap.batch_id == b.id).order_by(MigrationMap.id)
    maps, total = await paginate(session, stmt, maps_params)
    map_count = (
        await session.execute(
            select(func.count()).select_from(MigrationMap).where(MigrationMap.batch_id == b.id)
        )
    ).scalar_one()
    assert map_count == total
    return S.MigrationBatchDetail(
        **base.model_dump(),
        errors=[S.ImportError(**e) for e in b.errors],
        duplicates=[S.ImportDuplicate(**d) for d in b.duplicates],
        maps=Page[S.MigrationMap](
            items=[S.MigrationMap.model_validate(m) for m in maps],
            page=maps_params.page,
            size=maps_params.size,
            total=total,
        ),
    )
